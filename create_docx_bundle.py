from pathlib import Path
from docx import Document
import re

root = Path(r"c:\Users\MANASWINI\Downloads\mediassist_ai (1)")
output = root / 'mediassist_all_files.docx'

def sanitize_text(text: str) -> str:
    return re.sub(r'[\x00-\x08\x0B\x0C\x0E-\x1F\x7F]', '', text)

if output.exists():
    output.unlink()

print(f"Scanning files under {root}")
doc = Document()
doc.add_heading('MediAssist Project Files', level=1)
doc.add_paragraph('Generated from workspace files.')

count = 0
ignored_dirs = {'.git', '.venv', '.venv-1', '.gitignore', 'workspaceStorage'}
for path in sorted(root.rglob('*')):
    if not path.is_file():
        continue
    if any(part in ignored_dirs for part in path.parts):
        continue
    if path.suffix.lower() in {'.png', '.jpg', '.jpeg', '.gif', '.ico', '.bmp', '.pdf'}:
        continue
    try:
        raw = path.read_bytes()
        if b'\x00' in raw:
            print(f"Skipping binary-like file: {path}")
            continue
        try:
            text = raw.decode('utf-8')
        except UnicodeDecodeError:
            text = raw.decode('latin-1')
    except Exception as e:
        print(f"Could not read {path}: {e}")
        continue
    text = sanitize_text(text)
    rel = path.relative_to(root)
    doc.add_page_break()
    doc.add_heading(str(rel), level=2)
    for line in text.splitlines():
        doc.add_paragraph(line)
    count += 1

print(f"Added {count} files")
doc.save(output)
print(f"Saved {output}")
